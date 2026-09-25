"""Create the English report with the Codex bundled Python runtime."""
from pathlib import Path
import json, csv, shutil
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[2]; OUT=ROOT/'output/report'; AS=OUT/'assets'; BASE=ROOT/'results/report_20260925'
BLUE='184991'; INK='192638'; PALE='EEF3FA'
ev=json.loads((BASE/'reviewed_test/summary.json').read_text())
assert len(ev['results'])==2, 'Both checkpoint evaluations must finish before authoring.'
old,new=[r['metrics'][1] for r in ev['results']]
hist=json.loads((BASE/'historical_recomputed.json').read_text())
sim={k:json.loads((BASE/f'sim/{k}/check.json').read_text()) for k in ['static','range-sweep','weave']}
train=list(csv.DictReader(open(ROOT/'weights/boat_reviewed_20260917_results.csv')))
best=max(train,key=lambda r:float(r['metrics/mAP50-95(B)']))
doc=Document(); sec=doc.sections[0]
sec.page_width=Inches(8.27);sec.page_height=Inches(11.69)
sec.top_margin=Inches(.7);sec.bottom_margin=Inches(.67);sec.left_margin=sec.right_margin=Inches(.75)
sec.header_distance=Inches(.3);sec.footer_distance=Inches(.3)
styles=doc.styles
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption']:
 st=styles[name];st.font.name='Calibri';st.font.color.rgb=RGBColor.from_string('000000')
normal=styles['Normal'];normal.font.size=Pt(11);normal.paragraph_format.line_spacing=1.08;normal.paragraph_format.space_after=Pt(7)
for name,size in [('Title',32),('Subtitle',15),('Heading 1',22),('Heading 2',13),('Heading 3',11)]:
 styles[name].font.size=Pt(size);styles[name].font.bold=name!='Subtitle'
 styles[name].paragraph_format.space_after=Pt(9);styles[name].paragraph_format.space_before=Pt(12)
styles['Caption'].font.size=Pt(9);styles['Caption'].font.color.rgb=RGBColor.from_string('526477');styles['Caption'].font.italic=False;styles['Caption'].font.bold=False
styles['Caption'].paragraph_format.space_after=Pt(10)
styles['Caption'].paragraph_format.keep_with_next=False
for st in styles:
 for border in list(st.element.findall('.//' + qn('w:pBdr'))):border.getparent().remove(border)
header=sec.header.paragraphs[0];header.text='SHIP DETECTION    /    INTERNSHIP TECHNICAL REPORT';header.style='Caption'
for run in header.runs:run.font.color.rgb=RGBColor(0,0,0);run.font.size=Pt(8)
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=footer.add_run('Repository review  •  25 September 2026     |     ');r.font.size=Pt(8);r.font.color.rgb=RGBColor.from_string('526477')
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
doc.core_properties.title='Ship Detection and Distance Estimation'
doc.core_properties.subject='Repository analysis, measured detector comparison and basin simulation'
doc.core_properties.author='';doc.core_properties.keywords='ship detection, computer vision, UWB, simulation, internship'

def p(text,style=None):return doc.add_paragraph(text,style)
def h(text):return doc.add_heading(text,2)
def page(num,title):
 label=p(f'{num:02d}   /   TECHNICAL REVIEW','Caption');label.paragraph_format.page_break_before=True;label.paragraph_format.keep_with_next=True;doc.add_heading(title,1)
def picture(path,caption,width=6.65):
 para=doc.add_paragraph();para.paragraph_format.space_after=Pt(4);para.paragraph_format.keep_with_next=True
 para.add_run().add_picture(str(path),width=Inches(width));p(caption,'Caption')
def table(headers,rows,widths=None):
 t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 widths=widths or [6.65/len(headers)]*len(headers)
 for col,w in zip(t.columns,widths):col.width=Inches(w)
 for i,name in enumerate(headers):t.rows[0].cells[i].text=str(name)
 for row in rows:
  cells=t.add_row().cells
  for i,value in enumerate(row):cells[i].text=str(value)
 for ri,row in enumerate(t.rows):
  pr=row._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');pr.append(cant)
  if ri==0:
   rep=OxmlElement('w:tblHeader');pr.append(rep)
  for ci,cell in enumerate(row.cells):
   cell.width=Inches(widths[ci]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   tc=cell._tc.get_or_add_tcPr();sh=OxmlElement('w:shd');sh.set(qn('w:fill'),BLUE if ri==0 else ('FFFFFF' if ri%2 else PALE));tc.append(sh)
   mar=OxmlElement('w:tcMar')
   for name in ['top','left','bottom','right']:
    e=OxmlElement('w:'+name);e.set(qn('w:w'),'90');e.set(qn('w:type'),'dxa');mar.append(e)
   tc.append(mar);b=OxmlElement('w:tcBorders')
   for name in ['top','left','bottom','right']:
    e=OxmlElement('w:'+name);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');b.append(e)
   tc.append(b)
   for para in cell.paragraphs:
    para.paragraph_format.space_after=Pt(1);para.paragraph_format.line_spacing=1.05
    if ci>0:para.alignment=WD_ALIGN_PARAGRAPH.CENTER
    for run in para.runs:run.font.size=Pt(9.5);run.font.bold=ri==0;run.font.color.rgb=RGBColor.from_string('FFFFFF' if ri==0 else INK)
 spacer=p('');spacer.paragraph_format.space_after=Pt(3);spacer.paragraph_format.line_spacing=Pt(2);spacer.add_run().font.size=Pt(2)
 return t
def source(text):p(text,'Caption')
def code(text):
 para=p(text);para.paragraph_format.space_after=Pt(8)
 for r in para.runs:r.font.name='Courier New';r.font.size=Pt(8.5)
def percent(x):return f'{100*x:.1f}%'

# 1
p('TECHNICAL REPORT   /   SEPTEMBER 2026','Caption')
doc.add_heading('Ship Detection and\nDistance Estimation',0)
p('Repository review and experimental results','Subtitle')
p('An internship report on camera perception, reviewed training data and a simulated two vessel test basin.')
picture(ROOT/'results/track_test/track_montage.jpg','Figure 1. Archived basin tracking example. Green boxes are tracked targets; the image is a tracking illustration, not a detector accuracy result. Source: results/track_test/track_montage.jpg and summary.json.')
h('Main finding')
p('The repository provides a working research workflow for detecting and tracking model boats, estimating distance and testing telemetry and motion algorithms. The main remaining task is to establish reliable performance on new basin recordings. Data review has made the training inputs more traceable, while the simulator exposes failures that ordinary unit tests do not catch.')
p(f'A new comparison for this report used the same 934 reviewed test images for the default detector and the September 17 fine tune. Their F1 scores were {old["f1"]:.3f} and {new["f1"]:.3f}, respectively, at confidence 0.45 and box IoU 0.5. These are diagnostic results on correlated recordings, with historical training exposure still requiring review.')
table(['Evidence checked','Scope'],[['Detector comparison','934 images and 362 labelled boats per model'],['Software checks','223 Python tests and two JavaScript checks passed'],['Fresh simulations','Static, range sweep and weave; one fixed seed each']],[2.4,4.25])
source('Snapshot: local working tree on 25 September 2026. Git HEAD: 22d94ed02b87, with substantial uncommitted changes. Scope: repository implementation and available experimental evidence.')

# 2
page(1,'Project scope and system design')
p('The intended setting is a covered test basin with two instrumented model vessels. A camera on one vessel observes the other. The software must keep a target identity over time and turn image measurements into useful distance and motion estimates. Reflections, small targets, low capture rates and camera movement complicate that task.')
picture(AS/'architecture.png','Figure 2. Main data paths reconstructed from run.py, boatdet/pipeline.py, the telemetry scripts and the dataset tools. Optional inputs require their own calibration and validation.')
h('Processing a recording')
p('run.py accepts a recorded video, raw rig frames or a live camera. YOLO produces candidate boat boxes. The single target path supports CSRT and Kalman tracking; the multi target path uses ByteTrack style association with persistent identifiers. Detection, tracking, ranging and optional target motion analysis (TMA) have separate modules under boatdet/.')
p('Each processed recording can produce an annotated video, measurements in CSV and metadata with processing settings. The local viewer presents recordings and can start processing jobs. A separate review interface records annotation decisions in SQLite and exports versioned labels for training.')
h('Boundaries of the repository')
p('This is a collection of Python tools and a local web interface. The software that controls the physical rigs belongs to the separate hydrolink repository. The detector and range tools can run offline, whereas real telemetry collection depends on the laboratory network, synchronized capture and access to the instruments.')
p('Segmentation, AprilTag observations, attitude compensation and world frame motion analysis are optional. Their presence in the architecture does not establish that all inputs are available or accurate. In particular, the repository ships no maritime segmentation checkpoint, and the current documentation does not establish a validated own ship navigation reference for real TMA trials.')
source('Sources: README.md; HANDOFF.md; boatdet/pipeline.py; docs/TRACKING_SEGMENTATION.md.')

# 3
page(2,'Reviewed data and training splits')
p('The most recent reviewed source contains 5,016 frames. The September 17 materialization uses saved review decisions, checks image and label hashes, resolves recording identity where possible and holds unresolved records outside the active splits. Positive frames contain at least one boat box; the number of boxes can therefore exceed the number of positive frames.')
picture(AS/'dataset.png','Figure 3. Allocation of the 5,016 reviewed source frames. Held and excluded frames are outside the active train, validation and test datasets. Counts come directly from the frozen manifest.')
table(['Split','Frames','Positive','Background','Boxes'],[['Training','1,807','808','999','961'],['Validation','411','174','237','197'],['Test','934','336','598','362'],['Held','1,799','70','1,729','72'],['Excluded','65','2','63','2']],[1.65,1.25,1.25,1.3,1.2])
p('The active splits contain 3,152 images. Another 1,799 frames remain held, mainly because their recording identity is unresolved. Keeping those images out of training prevents uncertain session membership from silently weakening the evaluation split. The saved audit reports no structural issues for the active dataset and records image decoding and hashing as enabled.')
h('What this split can establish')
p('Whole recording sessions are assigned to one split. The manifest names three validation sessions and five test sessions. This is stronger provenance than a folder name alone, but it does not reconstruct every image used by older checkpoints. Reusing a previously trained model as a starting point also carries its historical exposure into the new experiment.')
p('The older colour dataset has different counts and annotation history. Its September 14 snapshot contains 2,466 training, 831 validation and 815 test directory images. Those totals should not be combined with the reviewed split or presented as one dataset.')
source('Sources: results/reviewed_yolo_20260917/dataset/manifest.json; audit.json; prepare_reviewed_training.py; docs/AUDIT_20260914.md.')

# 4
page(3,'Annotation quality affects the score')
p('The earlier audits found disagreements on identical images, including visible boats marked as background and boxes drawn around different extents of the same target. These differences affect both learning and evaluation: a correct boat location can fail an IoU threshold when the reference box follows a different convention.')
picture(ROOT/'results/manual_training_review_20260914/annotation_versions.png','Figure 4. Archived comparison of rig source labels in red and the then current local training labels in cyan. The image documents earlier label conflicts; it is not an assertion that these conflicts remain in the reviewed September 17 labels.',width=5.85)
p('In one duplicated colour recording, two annotation passes had a median box IoU of 0.481 across 159 identical frames. In the separate manual source comparison, 214 valid positive pairs had a median IoU of 0.552, and 74 pairs fell below 0.5. The groups and dates differ, so the two statistics describe separate checks.')
p('A practical annotation rule should specify the visible hull extent, treatment of reflections, partially occluded boats and secondary targets. Negative labels need a deliberate check that no boat is visible. The review interface and decision snapshot support this process, but structural validation alone cannot certify label completeness.')
source('Sources: docs/DATASET_REPAIR.md; docs/MANUAL_DATA_REVIEW_20260914.md; results/manual_training_review_20260914/annotation_versions.png.')

# 5
page(4,'Fresh detector comparison')
p('Both checkpoints were evaluated for this report on all 934 images in the reviewed test split, containing 362 labelled boat instances. The comparison uses CPU FP32, four PyTorch threads, image size 960 and non maximum suppression (NMS) IoU 0.7. Predictions are collected at confidence 0.15; the primary scores below use confidence 0.45 and matching intersection over union (IoU) 0.5.')
picture(AS/'detector_comparison.png','Figure 5. New paired checkpoint comparison on the reviewed test split. Precision measures how many predicted boxes match labels; recall measures how many labelled boats are found. F1 balances both.')
table(['Checkpoint','TP','FP','FN','Precision','Recall','F1'],[['Default v4',old['tp'],old['fp'],old['fn'],percent(old['precision']),percent(old['recall']),f'{old["f1"]:.3f}'],['Reviewed fine tune',new['tp'],new['fp'],new['fn'],percent(new['precision']),percent(new['recall']),f'{new["f1"]:.3f}']],[1.75,.6,.6,.6,1.05,1,.85])
source('TP: matched predictions. FP: unmatched predictions. FN: labelled boats missed by the model.')
delta=new['f1']-old['f1']
p(f'The fine tune changes F1 by {delta:+.3f}. It produces {new["tp"]} matched detections and {new["fp"]} unmatched predictions, compared with {old["tp"]} and {old["fp"]} for the default checkpoint. This comparison supports a decision about these images and settings; it does not establish performance across new outings, lighting conditions or different target vessels.')
if new['f1']<old['f1']:
 p('The result does not support replacing the default checkpoint with the reviewed fine tune. Reviewed labels are valuable, but this particular training run has not converted that improvement in provenance into a better score on the reviewed test images.')
else:
 p('The fine tune scores better on this diagnostic set. A promotion decision still needs a separate recording set with documented training exclusion and a review of errors by operating condition.')
h('Interpretation limits')
p('Frames from each recording are strongly correlated. Historical model exposure is not fully reconstructed, and labels can still omit targets or disagree on box extent. An unmatched prediction is a false positive against the supplied labels, not automatically proof that the image contains no vessel. The latest checkpoint was compared directly; no retraining or deployment was performed for this report.')
source('Fresh evidence: results/report_20260925/reviewed_test/summary.json and both per image prediction files. Checkpoints: weights/boat_v4_s_best.pt and weights/boat_reviewed_20260917_best.pt.')

# 6
page(5,'Earlier experiments and their limits')
p('The September 14 audit compared six checkpoints on the same saved images. Its newer diagnostic subset consists of 704 frames from seven September 10 recordings, with 168 labelled boats. The results below were recomputed for this report by summing the saved per recording counts and excluding the historical renkliTekneTekne recording.')
names={'boat_v4_s_best':'v4 default','boat_v4s_frozen_best':'v4 frozen','boat_v4s_frozen_color_best':'v4 colour','boat_v4s_frozen_v3_best':'v3 frozen','boat_v4s_frozen_v3_unfrozen_best':'v3 unfrozen','boat_v5_s_best':'v5'}
table(['Checkpoint','TP','FP','FN','Precision','Recall','F1'],[[names[r['model']],r['tp'],r['fp'],r['fn'],percent(r['precision']),percent(r['recall']),f'{r["f1"]:.3f}'] for r in hist],[1.75,.6,.6,.6,1.05,1,.85])
p('The default model leads this diagnostic subset by F1, whereas the validation ranking differs. That disagreement is a reason to inspect recording level errors before choosing a checkpoint. The small difference between the default and v3 unfrozen is not evidence of a consistent advantage over independent operating conditions.')
h('Native colour input')
p('The September 10 paired experiment used 245 aligned frames from one recording, with 47 labelled boats and 198 labelled negative frames. For the default v4 model, replacing an intermediate 640 by 360 resize with the decoded native colour input changed true positives from 22 to 23 while false positives remained 28. Recall rose from 46.8% to 48.9%, and F1 from 0.454 to 0.469.')
p('The native grayscale arm found no true positives for the default model at confidence 0.45. This is evidence against converting these colour trained checkpoints to grayscale at inference. It does not test a detector trained on grayscale imagery. The active training workflow therefore uses colour recordings.')
h('Duplicate suppression')
p('The audit also isolated global NMS on saved tiled predictions. In the full tiles arm, false positives fell from 515 to 237, while true positives fell from 365 to 357. Removing duplicate boxes reduced false alarms but also removed some valid matches. This tradeoff does not by itself justify enabling tiled detection for every recording.')
p('None of these older scores should be directly compared with the fresh 934 image experiment as though the dataset were unchanged. The annotation version, split composition and model history differ. Their value is in identifying specific problems and checking proposed changes under a controlled comparison.')
source('Sources: docs/EVALUATION.md; docs/AUDIT_20260914.md; results/audit_20260914/models/summary.json; results/report_20260925/historical_recomputed.json.')

# 7
page(6,'The reviewed fine tuning run')
p('The saved September 17 training log contains 17 epochs. The associated arguments request up to 40 epochs with patience 10, starting from the default v4 checkpoint. The report uses the CSV and arguments stored beside the released weight file, rather than treating an old local process status as proof that training is still active.')
picture(AS/'training.png','Figure 6. Saved validation metrics and box losses for the reviewed fine tune. The dotted line marks epoch 7, the highest mAP50-95 row in the saved log. Source: weights/boat_reviewed_20260917_results.csv.')
table(['Setting or result','Recorded value'],[['Optimizer and initial learning rate','AdamW; 0.0002'],['Image size and batch size','960; 16'],['Seed and requested budget','17; 40 epochs; patience 10'],['Best logged mAP50-95','0.20418 at epoch 7'],['mAP50 at the same epoch','0.47284'],['Precision and recall at that epoch','0.58301 and 0.43147'],['Final logged mAP50-95','0.15720 at epoch 17']],[3.6,3.05])
p('Training box loss continues to decrease while validation box loss rises after its early minimum. The pattern is consistent with overfitting or a mismatch between training and validation conditions. It does not identify the cause on its own. The saved metrics do show that running more epochs did not steadily improve validation quality.')
p('mAP50 summarizes the precision recall curve at box IoU 0.5. mAP50-95 averages this quantity across IoU thresholds from 0.5 to 0.95. These are different measurements from F1 at confidence 0.45, so the training graph cannot replace the paired checkpoint test on the previous pages.')
p('The next training experiment should keep the evaluation sessions fixed, record the exact source weight hash and data manifest, and change one main factor at a time. Reviewing missed boats and false alarms by recording is more informative than selecting a run from its lowest training loss.')
source('Sources: weights/boat_reviewed_20260917_args.yaml; weights/boat_reviewed_20260917_results.csv; fresh paired test in results/report_20260925/reviewed_test/.')

# 8
page(7,'Tracking and distance estimation')
h('Maintaining a target identity')
p('The multi object tracker first predicts existing tracks to the capture timestamp, then associates strong detections using overlap and appearance. Lower confidence detections may continue an existing track. Motion gates handle movement that is too large for overlap alone. The implementation uses image timing rather than assuming every recording has a fixed frame rate.')
p('An optional occlusion setting retains a confirmed vessel identity through a short absence. During that period, the displayed box is a prediction. The current pipeline does not sample the predicted background or an old water contact as a fresh range observation. CSV fields expose tracking status and measurement age so downstream analysis can distinguish observations from estimates.')
p('That distinction matters at the roughly 4 to 6 frames per second seen in archived rig recordings. A frame count alone corresponds to different real durations at different capture rates. Unit tests cover several occlusion cases, but the repository does not provide an independent real scene benchmark of identity switches or occlusion recovery rates.')
h('Choosing the range source')
table(['Available source','How range is obtained','Main dependency'],[['AprilTag pose','Metric position from a tag pose','A valid pose, not only an angle'],['Water contact','Camera ray intersected with water plane','Height, attitude and contact pixel'],['Monocular depth','Median positive depth in a target region','Model transfer and rig calibration']],[1.35,2.7,2.6])
p('The multi target pipeline gives priority to a valid tag pose, then to calibrated water plane geometry, then to monocular depth. The default depth model is Depth Anything V2 Metric Indoor Base. The repository starts with identity distance calibration; a fit obtained elsewhere should not be treated as a universal correction for this camera and basin.')
h('Why the waterline matters')
p('For a level camera above a flat water plane, horizontal range is approximately camera height divided by the tangent of the downward viewing angle. As that angle approaches zero near the horizon, a small pixel or pitch error produces a large range error. A reflection or a box bottom below the real waterline therefore biases the result directly.')
p('Optional segmentation can refine the water contact location, but no suitable checkpoint is bundled. The fallback contact remains dependent on detector box geometry. Course, speed and time to collision inherit the uncertainty of range, association and timing; a numeric overlay alone is not validation of those quantities.')
source('Sources: boatdet/pipeline.py; boatdet/depth.py; boatdet/geometry.py; boatdet/mot.py; docs/TRACKING_SEGMENTATION.md.')

# 9
page(8,'Telemetry and calibration')
p('The telemetry path converts ultra-wideband (UWB) ranges and inertial measurement unit (IMU) messages into per ship state histories. The producer solves position from surveyed anchors and combines tag geometry with attitude. These histories align camera observations with a reference frame for calibration and supervision.')
h('Reference frames must agree')
p('A camera angle is initially relative to the camera optical axis. Converting it into a world bearing requires the camera mounting angle and vessel heading on the same clock. Camera position also differs from the navigation reference point by a lever arm. Ignoring any of these terms can turn a calibration bias into apparent target motion.')
p('The archived tracking example contains 84 frames: 74 tracked, 3 lost and 7 marked off water. Its bearings are explicitly camera relative. The example demonstrates tracking and angle extraction but provides no independent range or course reference.')
h('Bearing uncertainty')
p('The earlier source video jitter study measured 3,534 residuals from 41 chains across eight recordings. Raw box centre jitter was 0.69 pixels RMS in the 640 by 360 working frame. The measurement removes smooth motion and therefore does not measure slow aspect dependent bias. A filtered track centre is also not a substitute for raw detector jitter.')
table(['Error budget term','Documented working assumption'],[['Pixel localisation','1.5 px, equivalent to about 0.14 degrees at f = 600 px'],['Vessel heading','1.0 degree'],['Camera mount','0.3 degrees'],['Combined bearing sigma','About 1.05 degrees by quadrature']],[2.2,4.45])
p('Heading dominates this example budget. Improving subpixel localisation cannot compensate for an unknown vessel heading or an incorrect camera mount. The numerical budget is a configuration example tied to focal length and assumed uncertainties; it is not a measured total error distribution for the current hardware.')
h('Remaining physical measurements')
p('The handoff records unresolved camera extrinsics, unavailable second vessel cameras and the absence of a validated navigation source for real TMA evaluation. Those are historical inspection findings, not a fresh hardware health check. A new capture should include measured camera height and mounting angles, timestamp offsets and an independent position or distance reference for the target.')
p('Fit calibration on training sessions and evaluate it on held out sessions. The synthetic results favour straight moving legs for heading calibration: a stationary hull can produce misleading course samples from velocity noise.')
source('Sources: rig_state_logger.py; rig_state_producer.py; rig_calibration.py; docs/TMA.md; HANDOFF.md; results/track_test/summary.json.')

# 10
page(9,'Basin simulation design')
p('Three scenarios were rerun for this report with seed 0 and the default simulator settings: static for 60 seconds, weave for 90 seconds and range sweep for 200 seconds. The simulator writes telemetry, camera observations and ground truth in the same file formats used by the real tools. All three runs are synthetic.')
picture(AS/'trajectories.png','Figure 7. True simulated trajectories in the basin anchor frame. Circles mark starts and crosses mark ends. Squares and triangles indicate anchor positions. The water boundary is an assumed rectangle, not a new physical survey.')
h('What is being tested')
p('The producer check feeds synthetic telemetry through the actual ShipState implementation at 10 Hz. The camera check evaluates noisy bearings, waterline range and the range hypothesis TMA bank. It uses true observer position and heading when constructing TMA observations, so it does not measure a complete sensor fused perception system with uncertain navigation.')
p('The simulated camera does not render images for YOLO. Detection is sampled from a probability model based on target box height. A high synthetic tracked fraction therefore says nothing about the trained detector recall. The pursuit scenario exists in the repository but was not run here, and its controller would be a truth based baseline rather than the deployed DQN.')
h('Assumptions that limit transfer')
p('Anchor surveys and several measured noise characteristics inform the twin. Hull thrust, drag and mass are not physically identified. Camera height, mounting angles and the water boundary include placeholders; the anchor selection rule is an approximation. The optional steep path NLOS bias was left off, as the repository documents a conflict between that hypothesis and rig observations.')
p('Each scenario has one seed in this report. The results are reproducible examples for inspecting algorithm behaviour, not confidence intervals across noise realizations. A broader simulation campaign would be needed to estimate failure rates.')
source('Sources: pool_sim.py; boatdet/twin.py; docs/SIMULATION.md; results/report_20260925/sim/*/session.json. Fresh runs executed on 25 September 2026; synthetic sensor timestamps use the simulator epoch.')

# 11
page(10,'Results from the fresh simulations')
table(['Metric','Static','Range sweep','Weave'],[
 ['Camera frames',sim['static']['camera']['frames'],sim['range-sweep']['camera']['frames'],sim['weave']['camera']['frames']],
 ['Frames with synthetic tracks',sim['static']['camera']['tracked'],sim['range-sweep']['camera']['tracked'],sim['weave']['camera']['tracked']],
 ['Target in view',sim['static']['camera']['target_in_view'],sim['range-sweep']['camera']['target_in_view'],sim['weave']['camera']['target_in_view']],
 ['Ship 1 median position error',*[f"{sim[k]['producer']['1']['position_error_m']['median']:.3f} m" for k in sim]],
 ['Ship 2 median position error',*[f"{sim[k]['producer']['2']['position_error_m']['median']:.3f} m" for k in sim]],
 ['Median waterline range error',*[percent(sim[k]['camera']['waterline_relative_range_error']['median']) for k in sim]],
 ['90% TMA interval coverage',*[percent(sim[k]['camera']['tma_interval_coverage']) for k in sim]],
 ['TMA converged updates',*[sim[k]['camera']['tma_status_counts'].get('converged',0) for k in sim]],
 ],[3.05,1.2,1.2,1.2])
p('Position error is the horizontal producer error against known truth. Waterline error is absolute relative range error on valid camera observations. Interval coverage is the fraction of scored TMA updates for which the reported range interval contains the truth. These denominators differ and should not be interpreted as one combined success rate.')
h('Position and heading')
p('Median producer position error is approximately 2 to 3 centimetres under the default synthetic range model. In the weave scenario, the 95th percentile is 4.9 cm for ship 1 and 6.4 cm for ship 2. Those small errors depend on the simulator assumptions; they do not verify the real firmware position output or actual multipath performance.')
p('The static heading calibration is unreliable: angular spreads are about 94 and 127 degrees, and both results fail the simulator check that mirrors the tool warning. During the weave, the recovered tag mounting offsets differ from truth by 0.08 and 1.87 degrees. Movement creates useful heading information, whereas station keeping can turn velocity noise into a misleading course estimate.')
h('Waterline range and TMA')
p('Median waterline relative error ranges from 7.1% in weave to 14.7% in the range sweep. These are single observation errors at the assumed camera height. They do not include validation of a trained waterline detector or a measured mount on the real vessel.')
p('Static and range sweep remain initializing or ambiguous and have 100% interval coverage. Their point range estimates are still poor; a wide interval containing the truth does not make its centre accurate. The weave is more concerning: it reports convergence on 19 updates and its overall coverage falls to 45.3%. The next page examines that failure separately.')
source('Fresh evidence: results/report_20260925/sim/static/check.json; sim/range-sweep/check.json; sim/weave/check.json. Default models, seed 0; no NLOS hypothesis enabled.')

# 12
page(11,'False confidence after a bearing gap')
p('In the weave run, the target is visible in 214 of 377 camera frames and produces 148 synthetic tracked observations. The TMA bank receives discontinuous bearings as the target leaves the field of view. Its declared confidence does not consistently match its range error.')
picture(AS/'tma_failure.png','Figure 8. Fresh weave run, seed 0. Blue is estimated range, black is truth, and the shaded band is the reported 90% range interval. Orange circles mark updates reported as converged. Gaps break the plotted lines; the y axis is limited to 65 m to show the failure region.')
p('The 19 converged updates have a median absolute relative range error of 194.7%. Across all 148 updates, the interval contains truth only 45.3% of the time. At the last update, the estimator is ambiguous at 40.2 m while the true range is 8.6 m; its interval of 25.1 to 54.4 m still excludes truth.')
h('Mechanism identified in the code')
p('BearingOnlyEKF.predict clamps the elapsed prediction step to the configured maximum of 2 seconds. After a longer observation gap, the timestamp advances to the new time while state propagation and process noise use the shorter interval. The filter can therefore retain more confidence than the unobserved motion supports. The September 24 simulation notes separately isolate this mechanism with perfect and continuous bearing comparisons.')
p('The fresh run reproduces the failure symptom; it is not a new controlled proof of the cause. Code inspection also shows that rig_calibration.py starts a new estimator after gaps over 2 seconds. The main pipeline drops estimators when track identities disappear, but it has no equivalent elapsed bearing gap check for every surviving identity.')
h('Recommended repair and verification')
p('Choose an explicit gap policy: predict over the complete elapsed duration with appropriate numerical steps, or reinitialize after a defined loss of observations. Then test a retained identity through a long gap and compare confidence coverage with the continuous observation case. The acceptance criterion should include absence of confidently wrong ranges, not merely completion without an exception.')
source('Sources: fresh weave_tma_trace.csv and sim/weave/check.json; boatdet/tma.py, BearingOnlyEKF.predict; boatdet/pipeline.py, _analyze_motion; rig_calibration.py, evaluate; docs/SIMULATION.md.')

# 13
page(12,'Verification and next work')
p('The current working tree passed 223 Python unit tests, both JavaScript regression scripts and the dependency consistency check during report preparation. The tests cover dataset handling, evaluation arithmetic, tracking, geometry, TMA, UWB and simulation components. Passing them establishes checked software behaviour; it does not establish a field accuracy or safety claim.')
table(['Check','Result'],[['Python unittest discovery','223 tests passed'],['Viewer panel data checks','Passed'],['Dataset review JavaScript checks','Passed'],['Python dependency consistency','No broken requirements'],['Fresh detector runs','Both checkpoints scored on all 934 test images'],['Fresh synthetic runs','Static, range sweep and weave completed']],[3.75,2.9])
h('Recommended order of work')
table(['Priority','Next task','Evidence needed to finish'],[['1','Correct TMA handling of observation gaps','Regression with retained identities, long gaps and interval coverage'],['2','Review detector failures by session','Misses and false alarms inspected on reviewed and newly captured recordings'],['3','Measure the camera and navigation setup','Height, mount, clock offsets and independent target reference'],['4','Run a controlled new detector experiment','Frozen manifests and a recording set excluded from all training'],['5','Validate real tracking and range','Identity continuity, range error and latency on synchronized recordings']],[.65,2.7,3.3])
p('The immediate software issue is the estimator confidence after a gap. For perception, the next useful evidence is a documented set of new colour recordings with consistent boxes and a target distance reference. That dataset would separate training improvement from familiarity with earlier scenes.')
p('No operating requirement for acceptable miss rate, false alarm rate, range error or processing latency is recorded here. Those limits should be agreed before a deployment decision. Reporting raw scores with their conditions is more useful at this stage than declaring the system ready.')
p('The repository now supports repeatable experiments and traceable label decisions. The remaining work is measurable: resolve the gap failure, test the detector on excluded sessions and validate distance and motion against physical references.')
source('Fresh verification logs are preserved under results/report_20260925/. Existing source files, trained weights and live rig settings were not changed for this report.')

# 14
page(13,'Reproducing the results')
p('Run commands from the repository root using its existing virtual environment. The report reflects a modified working tree on Git HEAD 22d94ed02b871afac1ce7636448a0063b404d703. Checking out that commit alone will not reproduce the current code or restore local datasets. A selected source hash manifest is saved with the new evidence.')
h('Detector comparison')
code('venv/bin/python evaluate_models.py \\\n  --data results/reviewed_yolo_20260917/dataset \\\n  --splits test \\\n  --weights weights/boat_v4_s_best.pt \\\n    weights/boat_reviewed_20260917_best.pt \\\n  --out results/report_20260925/reviewed_test')
p('The summary records settings, dataset hash and both weight hashes. Per image predictions support later rescoring. The saved reviewed audit is older evidence; this report did not repeat a full image hash audit of every source frame outside the evaluated split.')
h('Simulation and tests')
code('venv/bin/python pool_sim.py run --scenario weave --seed 0 \\\n  --out results/report_20260925/sim/weave --check\n\nvenv/bin/python -m unittest discover -s tests\nnode tests/test_panel_data.js\nnode tests/test_dataset_review.js\nvenv/bin/python -m pip check')
p('Use static and range-sweep in place of weave with separate output folders to repeat the other scenarios. The fresh results use default durations and no --nlos-deg option. Map and error figures are generated from the saved ground truth and check outputs.')
h('Evidence index')
table(['Report content','Primary repository evidence'],[['System design and limitations','README.md; HANDOFF.md; boatdet/pipeline.py'],['Data and label history','docs/DATASET_REPAIR.md; docs/MANUAL_DATA_REVIEW_20260914.md'],['Reviewed split','results/reviewed_yolo_20260917/dataset/manifest.json'],['Historical detector scores','results/audit_20260914/models/summary.json'],['Fine tune history','weights/boat_reviewed_20260917_results.csv and args.yaml'],['Fresh detector comparison','results/report_20260925/reviewed_test/'],['Fresh simulation results','results/report_20260925/sim/'],['TMA mechanisms and assumptions','boatdet/tma.py; docs/TMA.md; docs/SIMULATION.md']],[2.25,4.4])
source('Evidence convention: “fresh” identifies runs performed for this report. Earlier findings retain their source date. Simulated outcomes are labelled synthetic. Recommended work is separate from measured results.')

path=OUT/'Ship_Detection_Technical_Report.docx';doc.save(path)
print(path)
